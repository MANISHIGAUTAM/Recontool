#!/usr/bin/env python3
"""Build the Entity x Tool STATE MATRIX workbook (the "Excel database").

This is the human-readable / spreadsheet projection of the event-based v2
architecture (docs/arch-v2/). It merges:

  data/tool_inventory.json   (532 canonical extracted tools)
  data/flow_plan.json        (consumes/produces/method per tool)
  data/added_tools.json      (11 recommended additions)
  config/categories.yaml     (28-category registry)

and emits  data/recontool_state_matrix.xlsx  with these sheets:

  00_README              how to read the workbook + the model in one screen
  01_Categories          the 28 categories (stage, active, priority, chain_to)
  02_Tools_Master        one row per tool: category, I/O types, output files
  03_Routing_Matrix      tool x input-type grid (which entity types trigger a tool)
  04_Seed_RootDomain     the tools a fresh ROOT DOMAIN seeds (the minimum bootstrap)
  05_Output_File_Map     per tool -> typed output file(s), multi-capability split
  06_Entity_State_DB     LONG-format work-item table (the real DB model), seeded
  07_Entity_Tool_Matrix  WIDE example: example entities x every tool (PENDING/-)
  08_Eliminate_Redundant dedupe candidates (covered by a better tool)

Pure stdlib + openpyxl + pyyaml. Run:  python -m recontool.inventory.build_state_matrix
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent.parent.parent
INVENTORY = REPO / "data" / "tool_inventory.json"
FLOW = REPO / "data" / "flow_plan.json"
ADDED = REPO / "data" / "added_tools.json"
CATS = REPO / "config" / "categories.yaml"
OUT = REPO / "data" / "recontool_state_matrix.xlsx"

# ── canonical entity/input types (columns of the routing matrix) ─────────────
INPUT_TYPES = [
    "domain", "subdomain", "live_host", "url", "endpoint", "ip", "cidr", "asn",
    "cert", "service", "email", "username", "org", "keyword", "repo", "js_file",
    "bucket", "document", "apk", "coords", "cpe",
]

# Coarse flow_plan "consumes" tokens -> the canonical input types they accept.
CONSUMES_MAP = {
    "domain": ["domain", "subdomain"],
    "domain_or_email": ["domain", "subdomain", "email"],
    "domain_or_live_host": ["domain", "subdomain", "live_host", "url"],
    "live_host": ["live_host", "url", "subdomain"],
    "ip_or_asn": ["ip", "asn"],
    "ip_domain_service": ["ip", "domain", "subdomain", "service"],
    "username_or_org": ["username", "org"],
    "org_or_domain": ["org", "domain", "subdomain"],
    "org_or_keyword": ["org", "keyword"],
    "domain_ip_email": ["domain", "subdomain", "ip", "email"],
    "none": [],
    "file_or_url": ["document", "url", "js_file"],
    "org_or_coords": ["org", "coords"],
    "keyword_or_domain": ["keyword", "domain", "subdomain"],
    "subdomain": ["subdomain"],
    "url": ["url", "endpoint"],
    "ip": ["ip"],
    "asn": ["asn"],
    "keyword": ["keyword"],
    "repo": ["repo"],
    "js_file": ["js_file"],
    "cpe": ["cpe", "service"],
    "bucket": ["bucket"],
    "any": list(INPUT_TYPES),
}

# produced type -> (typed output filename stem, DB destination)
PRODUCES_FILE = {
    "subdomain": ("subdomains.txt", "asset(type=subdomain)"),
    "subdomain_candidates": ("subdomain_candidates.txt", "(unconfirmed -> resolver)"),
    "ip": ("ips.txt", "asset(type=ip)"),
    "cidr": ("cidrs.txt", "asset(type=cidr)"),
    "asn": ("asns.txt", "asset(type=asn)"),
    "url": ("urls.txt", "endpoint(url)"),
    "endpoint": ("endpoints.txt", "endpoint(url)"),
    "path": ("paths.txt", "endpoint(url)"),
    "email": ("emails.txt", "asset(type=email)"),
    "tech": ("tech.jsonl", "asset.tech[]"),
    "service": ("services.jsonl", "service(port,product)"),
    "host_port": ("services.jsonl", "service(port)"),
    "bucket": ("buckets.txt", "asset(type=bucket)"),
    "secret": ("secrets_redacted.jsonl", "finding(kind=secret,REDACTED)"),
    "finding": ("findings.jsonl", "finding"),
    "screenshot": ("screenshots/", "endpoint.evidence_ref"),
    "cert": ("certs.jsonl", "asset(type=cert)"),
    "tag": ("tags.jsonl", "asset.tags[]"),
    "param": ("params.txt", "endpoint.params[]"),
    "js_file": ("js/", "artifact(js)"),
    "resolver_list": ("resolvers.txt", "resource"),
    "any": ("mixed.jsonl", "(routed by type)"),
    "none": ("(raw only)", "(none)"),
}

# Tools whose ONE binary does several jobs -> each mode writes a DIFFERENT typed file.
# slug -> [{mode, consumes[], produces[], output_file}]
MULTI_CAPABILITY = {
    "ffuf": [
        {"mode": "dir",   "consumes": ["live_host", "url"], "produces": ["path"],      "output_file": "paths.txt"},
        {"mode": "vhost", "consumes": ["live_host"],        "produces": ["subdomain"], "output_file": "subdomains.txt"},
        {"mode": "param", "consumes": ["url"],              "produces": ["param"],     "output_file": "params.txt"},
    ],
    "gobuster": [
        {"mode": "dir",   "consumes": ["live_host", "url"], "produces": ["path"],      "output_file": "paths.txt"},
        {"mode": "dns",   "consumes": ["domain"],           "produces": ["subdomain"], "output_file": "subdomains.txt"},
        {"mode": "vhost", "consumes": ["live_host"],        "produces": ["subdomain"], "output_file": "subdomains.txt"},
        {"mode": "s3",    "consumes": ["keyword"],          "produces": ["bucket"],    "output_file": "buckets.txt"},
    ],
    "amass": [
        {"mode": "enum",  "consumes": ["domain"], "produces": ["subdomain"], "output_file": "subdomains.txt"},
        {"mode": "intel", "consumes": ["asn", "org", "cidr"], "produces": ["domain", "asn"], "output_file": "domains.txt"},
    ],
    "nmap": [
        {"mode": "sV",  "consumes": ["ip", "subdomain"], "produces": ["service"], "output_file": "services.jsonl"},
        {"mode": "nse", "consumes": ["service"],         "produces": ["finding"], "output_file": "findings.jsonl"},
    ],
    "katana": [
        {"mode": "crawl", "consumes": ["live_host", "url"], "produces": ["url"],     "output_file": "urls.txt"},
        {"mode": "jc",    "consumes": ["live_host"],        "produces": ["js_file"], "output_file": "js_files.txt"},
    ],
    "gau": [
        {"mode": "urls", "consumes": ["domain", "subdomain"], "produces": ["url"],     "output_file": "urls.txt"},
        {"mode": "js",   "consumes": ["domain", "subdomain"], "produces": ["js_file"], "output_file": "js_files.txt"},
    ],
    "dnsx": [
        {"mode": "resolve", "consumes": ["subdomain"], "produces": ["ip"],        "output_file": "ips.txt"},
        {"mode": "cname",   "consumes": ["subdomain"], "produces": ["subdomain"], "output_file": "subdomains.txt"},
    ],
    "tlsx": [
        {"mode": "tls", "consumes": ["live_host", "ip"], "produces": ["cert"],      "output_file": "certs.jsonl"},
        {"mode": "san", "consumes": ["live_host", "ip"], "produces": ["subdomain"], "output_file": "subdomains.txt"},
    ],
    "httpx": [
        {"mode": "probe", "consumes": ["subdomain", "ip"], "produces": ["live_host"], "output_file": "live_hosts.txt"},
        {"mode": "tech",  "consumes": ["subdomain", "ip"], "produces": ["tech"],      "output_file": "tech.jsonl"},
    ],
    "arjun": [
        {"mode": "param", "consumes": ["url", "endpoint"], "produces": ["param"], "output_file": "params.txt"},
    ],
}

# Redundant / dedupe candidates -> covered_by + verdict.
ELIMINATE = [
    ("altdns", "alterx / dnsgen / gotator", "Eliminate", "Legacy permutation engine; modern generators are faster + maintained."),
    ("aquatone", "gowitness / eyewitness", "Deprecate", "Deprecated; screenshotting fully covered. Keep only if its clustering UI is wanted."),
    ("waybackurls", "gau / waymore", "Deprecate", "gau aggregates Wayback+CC+OTX+urlscan; waymore adds responses. Redundant alone."),
    ("sublist3r", "subfinder / amass", "Eliminate", "Unmaintained; subfinder+amass cover all its sources and more."),
    ("knockpy", "subfinder / puredns", "Deprecate", "Slower, fewer sources; superseded."),
    ("dirb", "ffuf / feroxbuster / dirsearch", "Eliminate", "Old single-threaded dir buster; modern fuzzers dominate."),
    ("dirbuster", "ffuf / feroxbuster", "Eliminate", "GUI-era tool; superseded by fast fuzzers."),
    ("gobuster-dns", "subfinder / puredns", "Keep(low)", "Keep gobuster for dir/vhost/s3; DNS mode is redundant with puredns."),
    ("linkfinder", "jsluice / xnLinkFinder", "Deprecate", "jsluice is the modern JS endpoint/secret miner."),
    ("massdns", "puredns (wraps massdns)", "Keep(dep)", "Keep as the resolver engine, but drive via puredns not directly."),
    ("theHarvester", "subfinder + hunter + intel feeds", "Keep(low)", "Broad but shallow; keep for email/OSINT breadth, deprioritize for subs."),
]


def _load():
    inv = json.loads(INVENTORY.read_text())
    flow = json.loads(FLOW.read_text())
    added = json.loads(ADDED.read_text())
    import yaml
    cats = yaml.safe_load(CATS.read_text())
    return inv, flow, added, cats


def _cat_names(cats: dict) -> dict[int, str]:
    return {c["id"]: c["name"] for c in cats.get("categories", [])}


def _norm_inputs(consumes_token: str) -> list[str]:
    return CONSUMES_MAP.get(consumes_token, [])


def _build_tools(inv, flow, added):
    """Return a unified list of tool dicts with normalized I/O + output files."""
    flow_by_slug = {t["slug"]: t for t in flow["tools"]}
    inv_by_slug = {t["slug"]: t for t in inv["tools"]}
    tools = []

    for slug, ft in flow_by_slug.items():
        it = inv_by_slug.get(slug, {})
        consumes_token = ft.get("consumes", "domain")
        input_types = _norm_inputs(consumes_token)
        produces = [ft.get("produces", "none")]
        caps = MULTI_CAPABILITY.get(slug)
        if caps:
            input_types = sorted({t for c in caps for t in c["consumes"]} | set(input_types))
            produces = sorted({p for c in caps for p in c["produces"]})
        out_files = sorted({PRODUCES_FILE.get(p, (p + ".txt", "?"))[0] for p in produces})
        tools.append({
            "slug": slug,
            "name": ft.get("name", slug),
            "category": ft.get("primary_section"),
            "sections": ft.get("sections", []),
            "classification": ft.get("classification", "passive"),
            "method": ft.get("method", "?"),
            "web_only": ft.get("web_only", False),
            "real_contract": ft.get("real_contract", False),
            "requires_key": "key per config" in (ft.get("params") or ""),
            "consumes_token": consumes_token,
            "input_types": input_types,
            "produces": produces,
            "output_files": out_files,
            "capabilities": caps or [],
            "store_db": ft.get("store_db", ""),
            "source": "extracted",
        })

    for at in added["tools"]:
        input_types = sorted({t for t in at["consumes"]})
        # expand coarse 'live_host' etc. through the same map where applicable
        expanded = set()
        for c in at["consumes"]:
            expanded |= set(CONSUMES_MAP.get(c, [c]))
        input_types = sorted(expanded)
        produces = at["produces"]
        out_files = sorted({cap["output_file"] for cap in at["capabilities"]}) or \
            sorted({PRODUCES_FILE.get(p, (p + ".txt", "?"))[0] for p in produces})
        tools.append({
            "slug": at["slug"], "name": at["name"], "category": at["category"],
            "sections": [at["section"]], "classification": at["classification"],
            "method": at["integration"], "web_only": False, "real_contract": False,
            "requires_key": bool(at.get("requires_key")),
            "consumes_token": "+".join(at["consumes"]),
            "input_types": input_types, "produces": produces, "output_files": out_files,
            "capabilities": at["capabilities"], "store_db": "", "source": "added",
        })

    tools.sort(key=lambda t: (t["category"] or 99, t["slug"].lower()))
    return tools


# ── styling helpers ──────────────────────────────────────────────────────────
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2F5496")
ADD_FILL = PatternFill("solid", fgColor="E2EFDA")
PEND_FILL = PatternFill("solid", fgColor="FFF2CC")
NA_FILL = PatternFill("solid", fgColor="F2F2F2")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")


def _header(ws, headers, widths=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = HEAD
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
    ws.freeze_panes = "A2"
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def sheet_readme(wb):
    ws = wb.active
    ws.title = "00_README"
    ws.column_dimensions["A"].width = 120
    lines = [
        ("Recontool — Entity x Tool State Matrix (event-based architecture v2)", True),
        ("", False),
        ("WHAT THIS IS", True),
        ("A spreadsheet projection of the runtime 'database'. Each discovered KEY (primary key) is a row;", False),
        ("each tool is tracked as a work-item with a STATE (pending/in-progress/completed/skipped).", False),
        ("", False),
        ("PRIMARY KEYS (rows of the live DB)", True),
        ("  domain | subdomain | live_host | ip | cidr | asn | cert | url | endpoint(subdomain+path) | bucket | repo | email", False),
        ("  The dedupe_key (e.g. 'domain:api.grow.in', 'livehost:https://api.grow.in:443', 'ip:203.0.113.10') is the real PK.", False),
        ("", False),
        ("THE LOOP (how a key drives execution)", True),
        ("  1. A tool runs and writes TYPED output files (subdomains.txt, ips.txt, paths.txt, ...).", False),
        ("  2. The AI/normalizer reads those files, extracts new KEYS, dedupes, and upserts them as rows.", False),
        ("  3. For each new key, its TYPE decides which tool work-items become PENDING (see 03_Routing_Matrix).", False),
        ("  4. A worker pops the highest-priority PENDING (key,tool) item, marks IN_PROGRESS, runs it,", False),
        ("     writes typed output, marks COMPLETED. New keys re-enter at step 1 until convergence.", False),
        ("", False),
        ("HOW TO READ EACH SHEET", True),
        ("  01_Categories        the 28 recon categories (stage / active / priority / chains).", False),
        ("  02_Tools_Master      every tool: category, input types it consumes, output files it writes.", False),
        ("  03_Routing_Matrix    tool x input-type grid. 'Y' = this entity type makes the tool PENDING.", False),
        ("  04_Seed_RootDomain   the exact tools a fresh ROOT DOMAIN bootstraps (your minimum seed).", False),
        ("  05_Output_File_Map   per tool -> typed output file(s). Multi-mode tools (ffuf, gobuster,", False),
        ("                       katana, gau, dnsx, tlsx, httpx) split outputs by produced type.", False),
        ("  06_Entity_State_DB   the REAL DB model (long format): one row per (key, tool) work-item.", False),
        ("  07_Entity_Tool_Matrix WIDE demo: example keys x every tool (PENDING where applicable).", False),
        ("  08_Eliminate_Redundant  tools whose job a better tool already covers.", False),
        ("", False),
        ("PRE-BUILT vs DYNAMIC", True),
        ("  Pre-built (static knowledge): 01,02,03,04,05,08 — regenerated by", False),
        ("    `python -m recontool.inventory.build_state_matrix`.", False),
        ("  Dynamic (per scan): 06 / 07 are generated at runtime from the static routing + live discoveries.", False),
        ("  Seeded here with a worked 'grow.in' example so the model is concrete.", False),
    ]
    for r, (text, bold) in enumerate(lines, 1):
        c = ws.cell(r, 1, text)
        if bold:
            c.font = Font(bold=True, size=12 if r == 1 else 11,
                          color="2F5496" if r == 1 else "000000")


def sheet_categories(wb, cats):
    ws = wb.create_sheet("01_Categories")
    _header(ws, ["ID", "Name", "Stage", "Active", "Priority", "Input types",
                 "Output types", "Chains to", "Recursion triggers"],
            [5, 38, 10, 8, 9, 30, 30, 16, 34])
    for c in cats.get("categories", []):
        ws.append([
            c["id"], c["name"], c.get("stage"), "active" if c.get("active") else "passive",
            c.get("priority"), ", ".join(c.get("input_types", [])),
            ", ".join(c.get("output_types", [])),
            ", ".join(str(x) for x in c.get("chain_to", [])),
            ", ".join(c.get("recursion_triggers", [])),
        ])
    ws.auto_filter.ref = ws.dimensions


def sheet_tools_master(wb, tools, cat_names):
    ws = wb.create_sheet("02_Tools_Master")
    _header(ws, ["Slug", "Name", "Cat", "Category name", "Class", "Method",
                 "Needs key", "Multi-mode", "Consumes (input types)",
                 "Produces", "Output files", "Source"],
            [22, 24, 5, 30, 9, 8, 9, 10, 34, 22, 32, 9])
    for t in tools:
        row = [
            t["slug"], t["name"], t["category"], cat_names.get(t["category"], "?"),
            t["classification"], t["method"], "yes" if t["requires_key"] else "",
            "yes" if t["capabilities"] else "",
            ", ".join(t["input_types"]), ", ".join(t["produces"]),
            ", ".join(t["output_files"]), t["source"],
        ]
        ws.append(row)
        if t["source"] == "added":
            for c in range(1, len(row) + 1):
                ws.cell(ws.max_row, c).fill = ADD_FILL
    ws.auto_filter.ref = ws.dimensions


def sheet_routing(wb, tools):
    ws = wb.create_sheet("03_Routing_Matrix")
    headers = ["Slug", "Cat", "Class"] + INPUT_TYPES
    _header(ws, headers, [22, 5, 9] + [10] * len(INPUT_TYPES))
    for t in tools:
        row = [t["slug"], t["category"], t["classification"]]
        for it in INPUT_TYPES:
            row.append("Y" if it in t["input_types"] else "")
        ws.append(row)
        for i, it in enumerate(INPUT_TYPES):
            cell = ws.cell(ws.max_row, 4 + i)
            if cell.value == "Y":
                cell.alignment = CENTER
                cell.fill = PEND_FILL
    ws.auto_filter.ref = ws.dimensions
    # rotate the type headers for readability
    for i in range(len(INPUT_TYPES)):
        ws.cell(1, 4 + i).alignment = Alignment(textRotation=60, horizontal="center")


def sheet_seed_root(wb, tools):
    ws = wb.create_sheet("04_Seed_RootDomain")
    _header(ws, ["Order", "Slug", "Cat", "Class", "Method",
                 "Produces", "Output files", "Why it fires on a root domain"],
            [7, 22, 5, 9, 8, 22, 30, 50])
    seed = [t for t in tools if "domain" in t["input_types"]
            and t["category"] not in (18,)]  # exclude reference indexes
    # cheapest/high-signal first: discovery cats before enrich
    order_pri = {1: 1, 2: 0, 3: 2, 17: 3, 4: 4, 6: 5, 7: 6, 9: 7, 10: 8,
                 14: 9, 19: 9, 5: 10, 24: 11}
    seed.sort(key=lambda t: (order_pri.get(t["category"], 50), t["slug"].lower()))
    for i, t in enumerate(seed, 1):
        ws.append([
            i, t["slug"], t["category"], t["classification"], t["method"],
            ", ".join(t["produces"]), ", ".join(t["output_files"]),
            f"consumes '{t['consumes_token']}' -> accepts a root domain as seed",
        ])
    ws.auto_filter.ref = ws.dimensions
    ws.cell(1, 1).comment = None


def sheet_output_map(wb, tools):
    ws = wb.create_sheet("05_Output_File_Map")
    _header(ws, ["Slug", "Cat", "Mode/Capability", "Consumes", "Produces",
                 "Typed output file", "DB destination"],
            [22, 5, 16, 26, 16, 26, 34])
    for t in tools:
        if t["capabilities"]:
            for cap in t["capabilities"]:
                produces = cap["produces"]
                dbdest = ", ".join(sorted({PRODUCES_FILE.get(p, ("", "?"))[1] for p in produces}))
                ws.append([t["slug"], t["category"], cap.get("mode", ""),
                           ", ".join(cap.get("consumes", t["input_types"])),
                           ", ".join(produces), cap["output_file"], dbdest])
        else:
            for p in t["produces"]:
                fstem, dbdest = PRODUCES_FILE.get(p, (p + ".txt", "?"))
                ws.append([t["slug"], t["category"], "(single)",
                           t["consumes_token"], p, fstem, dbdest])
    ws.auto_filter.ref = ws.dimensions


# entity_type -> which input types it "satisfies" when applying tools
ENTITY_SATISFIES = {
    "domain": {"domain"},
    "subdomain": {"subdomain"},
    "live_host": {"live_host", "subdomain", "url"},
    "ip": {"ip"},
    "asn": {"asn"},
    "endpoint": {"endpoint", "url"},
}

EXAMPLE_ENTITIES = [
    ("domain:grow.in", "domain", "grow.in", "", "seed", 0, "", "yes"),
    ("domain:api.grow.in", "subdomain", "api.grow.in", "domain:grow.in", "subfinder", 1, "unknown", "yes"),
    ("livehost:https://app.grow.in:443", "live_host", "https://app.grow.in:443",
     "domain:app.grow.in", "httpx", 1, "yes", "yes"),
    ("ip:203.0.113.10", "ip", "203.0.113.10", "domain:api.grow.in", "dnsx", 1, "n/a", "yes"),
    ("asn:AS13335", "asn", "AS13335", "ip:203.0.113.10", "asnmap", 2, "n/a", "yes"),
    ("ep:app.grow.in:GET:/admin", "endpoint", "https://app.grow.in/admin",
     "livehost:https://app.grow.in:443", "ffuf", 2, "yes", "yes"),
]


def sheet_state_db(wb, tools):
    """LONG format — the real DB model (one row per work-item). Seeded sample."""
    ws = wb.create_sheet("06_Entity_State_DB")
    _header(ws, ["key (PK)", "entity_type", "value", "parent_key", "depth",
                 "discovered_by", "is_live", "in_scope", "tool", "tool_cat",
                 "status", "priority", "output_file", "completed_at"],
            [30, 12, 30, 26, 6, 14, 8, 9, 18, 8, 12, 8, 22, 18])
    # seed: for grow.in (domain) expand a representative set of tools as PENDING,
    # show a couple COMPLETED to illustrate the lifecycle.
    sample_completed = {"subfinder", "crtsh", "amass", "asnmap"}
    domain_tools = [t for t in tools if "domain" in t["input_types"]
                    and t["category"] in (1, 2, 3, 6, 17)]
    domain_tools = sorted(domain_tools, key=lambda t: (t["category"], t["slug"]))[:16]
    pk, etype, val = EXAMPLE_ENTITIES[0][:3]
    for t in domain_tools:
        done = t["slug"] in sample_completed
        ws.append([
            pk, etype, val, "", 0, "seed", "", "yes", t["slug"], t["category"],
            "completed" if done else "pending",
            {1: 14, 2: 10, 3: 16, 6: 20, 17: 18}.get(t["category"], 50),
            (t["output_files"][0] if t["output_files"] else ""),
            "2026-06-10T12:00:00Z" if done else "",
        ])
    # a couple of derived keys to show recursion
    ws.append(["domain:api.grow.in", "subdomain", "api.grow.in", "domain:grow.in", 1,
               "subfinder", "unknown", "yes", "dnsx", 2, "pending", 10, "ips.txt", ""])
    ws.append(["ip:203.0.113.10", "ip", "203.0.113.10", "domain:api.grow.in", 1,
               "dnsx", "n/a", "yes", "naabu", 21, "pending", 50, "services.jsonl", ""])
    ws.append(["livehost:https://app.grow.in:443", "live_host", "https://app.grow.in:443",
               "domain:app.grow.in", 1, "httpx", "yes", "yes", "ffuf", 20, "pending", 60, "paths.txt", ""])
    for r in range(2, ws.max_row + 1):
        st = ws.cell(r, 11).value
        if st == "pending":
            ws.cell(r, 11).fill = PEND_FILL
    ws.auto_filter.ref = ws.dimensions


def sheet_wide_matrix(wb, tools):
    """WIDE demo — example entities (rows) x every tool (cols)."""
    ws = wb.create_sheet("07_Entity_Tool_Matrix")
    headers = ["key (PK)", "entity_type", "value", "depth", "is_live"] + [t["slug"] for t in tools]
    _header(ws, headers, [30, 12, 28, 6, 8] + [4] * len(tools))
    for ent in EXAMPLE_ENTITIES:
        pk, etype, val, parent, disc, depth, live, scope = ent
        satisfies = ENTITY_SATISFIES.get(etype, {etype})
        row = [pk, etype, val, depth, live]
        for t in tools:
            applicable = bool(satisfies & set(t["input_types"]))
            row.append("PEND" if applicable else "")
        ws.append(row)
        for i, t in enumerate(tools):
            cell = ws.cell(ws.max_row, 6 + i)
            if cell.value == "PEND":
                cell.fill = PEND_FILL
                cell.alignment = CENTER
            else:
                cell.fill = NA_FILL
    for i in range(len(tools)):
        ws.cell(1, 6 + i).alignment = Alignment(textRotation=90, horizontal="center")
    ws.freeze_panes = "F2"


def sheet_eliminate(wb):
    ws = wb.create_sheet("08_Eliminate_Redundant")
    _header(ws, ["Tool", "Covered by", "Verdict", "Rationale"], [18, 34, 12, 80])
    for tool, by, verdict, why in ELIMINATE:
        ws.append([tool, by, verdict, why])
        ws.cell(ws.max_row, 4).alignment = WRAP
    ws.auto_filter.ref = ws.dimensions


def main():
    inv, flow, added, cats = _load()
    cat_names = _cat_names(cats)
    tools = _build_tools(inv, flow, added)

    wb = Workbook()
    sheet_readme(wb)
    sheet_categories(wb, cats)
    sheet_tools_master(wb, tools, cat_names)
    sheet_routing(wb, tools)
    sheet_seed_root(wb, tools)
    sheet_output_map(wb, tools)
    sheet_state_db(wb, tools)
    sheet_wide_matrix(wb, tools)
    sheet_eliminate(wb)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}  ({len(tools)} tools, "
          f"{len(added['tools'])} added, {len(cats['categories'])} categories)")


if __name__ == "__main__":
    main()
